"""
文件解析层：处理上传的项目文件（zip/docx/pdf/txt）
职责：提取文本内容 + 识别核心文件，为 AI 分析做准备
"""
import os
import zipfile
import tempfile

# 代码文件白名单——只保留这些扩展名的文件
CODE_EXTENSIONS = {
    '.java', '.py', '.js', '.jsx', '.ts', '.tsx', '.vue',
    '.html', '.css', '.xml', '.json', '.yml', '.yaml',
    '.sql', '.md', '.txt', '.gradle', '.kt', '.go', '.rs',
    '.c', '.cpp', '.h', '.cs', '.php', '.rb', '.swift',
}

# 跳过这些目录和文件
SKIP_DIRS = {
    'node_modules', '.git', '__pycache__', 'venv', 'env',
    '.idea', '.vscode', 'dist', 'build', 'target', '.gradle',
    'vendor', 'bin', 'obj', '.next', '.nuxt', 'coverage',
}

SKIP_FILES = {
    '.DS_Store', 'Thumbs.db', '.gitignore', '.gitattributes',
    'package-lock.json', 'yarn.lock', 'pnpm-lock.yaml',
}

# 核心文件标识——优先发送给 AI 分析
IMPORTANT_FILES = {
    'README.md', 'readme.md', 'README',
    'pom.xml', 'package.json', 'requirements.txt', 'Pipfile',
    'build.gradle', 'settings.gradle', 'Dockerfile', 'docker-compose.yml',
    'application.yml', 'application.properties', 'application.yaml',
}

# 最大文本长度（字符数），超出则截断
MAX_TOTAL_TEXT = 500_000
MAX_SINGLE_FILE = 30_000


def parse_uploaded_file(file_content: bytes, filename: str) -> dict:
    """
    主入口：根据文件类型分发到不同的解析器
    返回：{ "dir_structure": str, "files": [{"path": str, "content": str}], "type": str }
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == '.zip':
        return _parse_zip(file_content)
    elif ext == '.docx':
        return _parse_docx(file_content)
    elif ext == '.pdf':
        return _parse_pdf(file_content)
    elif ext in ('.txt', '.md'):
        text = file_content.decode('utf-8', errors='ignore')
        return {
            "dir_structure": filename,
            "files": [{"path": filename, "content": text[:MAX_SINGLE_FILE]}],
            "type": "text",
        }
    else:
        raise ValueError(f"不支持的文件类型: {ext}，请上传 .zip / .docx / .pdf / .txt / .md 文件")


def _parse_zip(file_content: bytes) -> dict:
    """解析 zip 压缩包：提取目录结构 + 核心代码文件"""
    with tempfile.TemporaryDirectory() as tmpdir:
        zip_path = os.path.join(tmpdir, 'upload.zip')
        with open(zip_path, 'wb') as f:
            f.write(file_content)

        with zipfile.ZipFile(zip_path, 'r') as zf:
            zf.extractall(tmpdir)

        # 找解压后的根目录（可能有一层包裹目录）
        extracted = [d for d in os.listdir(tmpdir) if d != 'upload.zip']
        if len(extracted) == 1 and os.path.isdir(os.path.join(tmpdir, extracted[0])):
            root = os.path.join(tmpdir, extracted[0])
        else:
            root = tmpdir

        # 构建目录结构
        dir_lines = []
        files = []
        total_chars = 0

        for dirpath, dirnames, filenames in os.walk(root):
            # 过滤掉跳过的目录
            dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS]

            # 计算相对路径
            rel_dir = os.path.relpath(dirpath, root)
            if rel_dir == '.':
                rel_dir = ''
            depth = rel_dir.count(os.sep) + 1 if rel_dir else 0
            indent = '  ' * depth

            if rel_dir:
                dir_lines.append(f"{indent}{os.path.basename(dirpath)}/")

            for fname in sorted(filenames):
                if fname in SKIP_FILES:
                    continue
                ext = os.path.splitext(fname)[1].lower()
                if ext not in CODE_EXTENSIONS:
                    continue

                filepath = os.path.join(dirpath, fname)
                rel_path = os.path.relpath(filepath, root)
                dir_lines.append(f"{indent}  {fname}")

                # 读取文件内容
                try:
                    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                        content = f.read()

                    # 单文件截断
                    if len(content) > MAX_SINGLE_FILE:
                        content = content[:MAX_SINGLE_FILE] + "\n... (文件过长，已截断)"

                    # 总量控制
                    if total_chars + len(content) > MAX_TOTAL_TEXT:
                        break

                    files.append({
                        "path": rel_path,
                        "content": content,
                        "is_important": fname in IMPORTANT_FILES or _is_important_path(rel_path),
                    })
                    total_chars += len(content)
                except Exception:
                    continue

            if total_chars >= MAX_TOTAL_TEXT:
                break

        return {
            "dir_structure": '\n'.join(dir_lines[:500]),  # 目录结构最多500行
            "files": files,
            "type": "project",
        }


def _parse_docx(file_content: bytes) -> dict:
    """解析 Word 文档"""
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("需要安装 python-docx 库。运行: pip install python-docx")

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        tmp.write(file_content)
        tmp_path = tmp.name

    try:
        doc = Document(tmp_path)
        text = '\n'.join(para.text for para in doc.paragraphs if para.text.strip())

        # 也提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells)
                if row_text.strip(' |'):
                    text += f'\n{row_text}'

        return {
            "dir_structure": "Word 文档",
            "files": [{"path": "document.docx", "content": text[:MAX_SINGLE_FILE]}],
            "type": "document",
        }
    finally:
        os.unlink(tmp_path)


def _parse_pdf(file_content: bytes) -> dict:
    """解析 PDF 文档（仅支持文字型 PDF）"""
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        raise RuntimeError("需要安装 PyPDF2 库。运行: pip install PyPDF2")

    with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
        tmp.write(file_content)
        tmp_path = tmp.name

    try:
        reader = PdfReader(tmp_path)
        text = ''
        for page in reader.pages:
            page_text = page.extract_text() or ''
            text += page_text + '\n'

        if not text.strip():
            raise ValueError("PDF 文件未提取到文字内容，可能是扫描版 PDF，建议转换为 Word 格式后重试")

        return {
            "dir_structure": "PDF 文档",
            "files": [{"path": "document.pdf", "content": text[:MAX_SINGLE_FILE]}],
            "type": "document",
        }
    finally:
        os.unlink(tmp_path)


def _is_important_path(rel_path: str) -> bool:
    """判断文件路径是否是核心模块（Controller/Service/Model 等）"""
    path_lower = rel_path.lower()
    important_keywords = [
        'controller', 'service', 'repository', 'dao', 'model',
        'entity', 'config', 'main', 'app', 'router', 'handler',
        'middleware', 'schema', 'views', 'api',
    ]
    return any(kw in path_lower for kw in important_keywords)


def filter_core_files(files: list[dict], max_files: int = 8) -> list[dict]:
    """
    从所有文件中筛选核心文件，优先发送 important 文件 + 重要路径文件
    用于分层分析的第二层
    """
    # 按重要程度排序：IMPORTANT_FILES > 重要路径 > 其他
    def sort_key(f):
        name = os.path.basename(f['path'])
        if name in IMPORTANT_FILES:
            return (0, f['path'])
        if f.get('is_important'):
            return (1, f['path'])
        return (2, f['path'])

    sorted_files = sorted(files, key=sort_key)
    return sorted_files[:max_files]
